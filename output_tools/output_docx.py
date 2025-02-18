# 搞一个输出文档的东西，先把形似的问题解决了，然后再把别的逐渐往上加。
import docx
from text_transfer.text_transfer import *

class output_docx():
    def __init__(self):
        # 这个的说法是把各种需要存的东西都弄进来，然后存着，然后最后一声令下输出出去。这部分应该是无益于发文章的所以差不多得了
        self.text_dict = {} 
        self.file_name = "demo1.docx"
        self.template_name = "templates/docx_template.docx"
        self.document = docx.Document(self.template_name)
        self.text_transfer = text_transfer()
        pass

    def save_file(self):
        self.arrange_docx()
        full_name = r"auto_test/" + self.file_name
        self.document.save(full_name)
        pass

    def set_heading(self,heading_str:str):
        self.text_dict["heading"] = heading_str
    
    def set_jieguo(self,jieguo_dict:dict):
        self.text_dict["jieguo"] = jieguo_dict

    def set_context(self, context_str:str):
        self.text_dict["context"] = context_str
    
    def set_submission(self,submission_list:list):
        self.text_dict["submission_list"] = submission_list
    
    def set_ufunc_json(self, ufunc_json:dict):
        if not("ufunc_json_list" in self.text_dict):
            self.text_dict["ufunc_json_list"] = []
        self.text_dict["ufunc_json_list"].append(ufunc_json)

    def arrange_docx(self):
        #开始之前得先删除一下，不然就不太对了。
        self.document = docx.Document(self.template_name)

        self.document.add_heading(self.text_dict["heading"], level=1)

        self.document.add_heading("一、推演背景", level=2)

        self.document.add_paragraph(self.text_dict["context"])

        self.document.add_heading("二、子任务信息", level=2)

        # self.document.add_paragraph(self.text_dict["jieguo"])
        self.arrange_submission()

        self.document.add_heading("三、结论", level=2)

        self.document.add_paragraph("以上，本方案共包含"+str(len(self.text_dict["submission_list"]))+"个子任务。不难看出，在仿真推演开始前，方案智能生成分系统成功完成了子任务序列生成，形成了作战方案，保证了电子对抗推演的顺利开展。")

    def arrange_submission(self):
        # print("output_docx.arrange_submission: unfinished yet")
        for i in range(len(self.text_dict["submission_list"])):
            

            submission_single = self.text_dict["submission_list"][i]

            if submission_single.force_arrange == "装甲车等其他地面力量":
                # 于是这个追加一下干扰的。
                flag_ganrao = True
            else:
                flag_ganrao = False

            self.document.add_heading(submission_single.id_str, level=3)

            # 然后参与的单位和时间啥的写一下
            str_single = "在第" + str(submission_single.time_arrange[0]) + "帧到第" +str(submission_single.time_arrange[1]) + "帧期间，辅助决策系统为" + submission_single.force_arrange + "分配了一个任务，命令其" + submission_single.type_str+"，具体出击方向为" + submission_single.config_json["出击方向"] +"，解算得到预定任务范围"+ str(submission_single.space_arrange)+"。"

            self.document.add_paragraph(str_single)

            # 然后决策依据写一下。
            state_shuofa_str = self.text_dict["jieguo"]["state_str_list"][i] # 这个就是那一堆json的，原则上还得再转化一下。
            state_shuofa_str = self.text_transfer.cut_from_str(state_shuofa_str,"**","**",model="infinite")
            state_shuofa_str = self.text_transfer.clean_the_str(state_shuofa_str)
            self.document.add_paragraph(state_shuofa_str)

            # 然后是utility函数的那个
            try:
                utilitys_shuofa_str = self.text_dict["ufunc_json_list"][i]["explanation"]
            except:
                utilitys_shuofa_str = ""
            self.document.add_paragraph(utilitys_shuofa_str)

            if flag_ganrao:
                # 那就是这里要指定一下干扰的说法。
                self.get_dianci_str(i)

    def get_dianci_str(self, i):
        # 专门来一个，生成电子对抗相关内容的.
        submission_single = self.text_dict["submission_list"][i]

        id_str = "电子对抗" + str(i)

        self.document.add_heading(id_str, level=3)

        # 然后参与的单位和时间啥的写一下

        model_selected,submodel_selected = self.text_transfer.generate_ECM_model()
        
        str_single = "在第" + str(submission_single.time_arrange[0]) + "帧到第" +str(submission_single.time_arrange[1]) + "帧期间，辅助决策系统为" + "我方电子干扰车" + "分配了一个任务，命令其执行" + model_selected + "。"

        self.document.add_paragraph(str_single)

        # 然后也是姑且一些说法。
        str_single = "具体地，结合当前已生成和子任务和推定敌方态势信息，我方电子干扰车需要完成" + model_selected
        state_shuofa_str = self.text_dict["jieguo"]["state_str_list"][i]
        state_shuofa_str = self.text_transfer.cut_from_str(state_shuofa_str,"**","**",model="infinite")
        state_shuofa_str = self.text_transfer.clean_the_str(state_shuofa_str)
        str_single += state_shuofa_str 
        self.document.add_paragraph(str_single)


